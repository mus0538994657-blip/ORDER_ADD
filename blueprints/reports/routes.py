import io
from datetime import datetime
from flask import (
    Blueprint, render_template, request, send_file, current_app,
)
from flask_login import login_required
from extensions import db
from models.order import Order
from models.category import Category

reports_bp = Blueprint('reports_bp', __name__, url_prefix='/reports')


def _apply_filters(query, date_from_str, date_to_str, status_filter, category_filter):
    """يُطبّق الفلاتر على استعلام الطلبات."""
    try:
        if date_from_str:
            query = query.filter(
                Order.created_at >= datetime.strptime(date_from_str, '%Y-%m-%d')
            )
        if date_to_str:
            dt = datetime.strptime(date_to_str, '%Y-%m-%d').replace(
                hour=23, minute=59, second=59
            )
            query = query.filter(Order.created_at <= dt)
    except ValueError:
        pass
    if status_filter:
        query = query.filter_by(status=status_filter)
    if category_filter:
        query = query.filter_by(category_id=category_filter)
    return query


@reports_bp.route('/')
@login_required
def index():
    date_from_str = request.args.get('date_from', '')
    date_to_str = request.args.get('date_to', '')
    status_filter = request.args.get('status', '')
    category_filter = request.args.get('category', '', type=int)

    query = _apply_filters(
        Order.query, date_from_str, date_to_str, status_filter, category_filter
    )
    orders = query.order_by(Order.created_at.desc()).all()

    total_revenue = sum(o.final_price for o in orders)

    stats_by_status = {}
    for s in Order.STATUSES:
        filtered = [o for o in orders if o.status == s]
        stats_by_status[s] = {
            'count': len(filtered),
            'revenue': sum(o.final_price for o in filtered),
            'color': Order.STATUS_COLORS.get(s, 'secondary'),
        }

    stats_by_category = {}
    for o in orders:
        cat_name = o.category.name if o.category else 'غير محدد'
        if cat_name not in stats_by_category:
            stats_by_category[cat_name] = {'count': 0, 'revenue': 0.0}
        stats_by_category[cat_name]['count'] += 1
        stats_by_category[cat_name]['revenue'] += o.final_price

    categories = Category.query.order_by(Category.name).all()

    return render_template(
        'reports/index.html',
        orders=orders,
        total_revenue=total_revenue,
        stats_by_status=stats_by_status,
        stats_by_category=stats_by_category,
        categories=categories,
        statuses=Order.STATUSES,
        date_from=date_from_str,
        date_to=date_to_str,
        status_filter=status_filter,
        category_filter=category_filter,
    )


# ---------------------------------------------------------------------------
# تصدير Excel
# ---------------------------------------------------------------------------

@reports_bp.route('/export/excel')
@login_required
def export_excel():
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    date_from_str = request.args.get('date_from', '')
    date_to_str = request.args.get('date_to', '')
    status_filter = request.args.get('status', '')
    category_filter = request.args.get('category', '', type=int)

    orders = _apply_filters(
        Order.query, date_from_str, date_to_str, status_filter, category_filter
    ).order_by(Order.created_at.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'تقرير الطلبات'
    ws.sheet_view.rightToLeft = True

    hdr_fill = PatternFill('solid', fgColor='1B4F72')
    alt_fill = PatternFill('solid', fgColor='EBF5FB')
    thin = Side(style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center')

    # عنوان التقرير
    ws.merge_cells('A1:L1')
    ws['A1'] = f'تقرير الطلبات — {datetime.utcnow().strftime("%Y-%m-%d")}'
    ws['A1'].font = Font(bold=True, size=14, color='1B4F72')
    ws['A1'].alignment = center
    ws.row_dimensions[1].height = 28

    headers = [
        'رقم الطلب', 'العميل', 'الهاتف', 'المركبة',
        'الخدمة', 'السعر', 'الخصم', 'الإجمالي',
        'الحالة', 'ملاحظات', 'تاريخ الإنشاء',
    ]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.font = Font(bold=True, color='FFFFFF', size=11)
        cell.fill = hdr_fill
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[2].height = 20

    for r, order in enumerate(orders, 3):
        row_data = [
            order.order_number, order.customer_name,
            order.customer.phone if order.customer else '',
            order.vehicle_info,
            order.category.name if order.category else '',
            order.price, order.discount, order.final_price,
            order.status, order.notes or '',
            order.created_at.strftime('%Y-%m-%d') if order.created_at else '',
        ]
        for c, val in enumerate(row_data, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
            if r % 2 == 0:
                cell.fill = alt_fill

    total_row = len(orders) + 3
    ws.cell(row=total_row, column=1, value='الإجمالي').font = Font(bold=True)
    ws.cell(row=total_row, column=8,
            value=sum(o.final_price for o in orders)).font = Font(bold=True)

    for i, w in enumerate([15, 20, 15, 22, 18, 10, 10, 10, 16, 25, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f'report_{datetime.utcnow().strftime("%Y%m%d_%H%M%S")}.xlsx'
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename,
    )


# ---------------------------------------------------------------------------
# تصدير PDF
# ---------------------------------------------------------------------------

@reports_bp.route('/export/pdf')
@login_required
def export_pdf():
    date_from_str = request.args.get('date_from', '')
    date_to_str = request.args.get('date_to', '')
    status_filter = request.args.get('status', '')
    category_filter = request.args.get('category', '', type=int)

    orders = _apply_filters(
        Order.query, date_from_str, date_to_str, status_filter, category_filter
    ).order_by(Order.created_at.desc()).all()

    filters = {
        'date_from': date_from_str,
        'date_to': date_to_str,
        'status': status_filter,
    }

    from utils.pdf import generate_orders_pdf
    workshop_name = current_app.config.get('WORKSHOP_NAME', 'ورشة الزجاج')
    buf = generate_orders_pdf(orders, filters, workshop_name=workshop_name)
    filename = f'report_{datetime.utcnow().strftime("%Y%m%d_%H%M%S")}.pdf'
    return send_file(buf, mimetype='application/pdf',
                     as_attachment=True, download_name=filename)

